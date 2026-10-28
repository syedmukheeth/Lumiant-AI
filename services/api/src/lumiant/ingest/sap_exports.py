"""SAP list-export intake (M8): MB52-style warehouse stock list + material master valuation list → DuckDB.

Deliberately strict (docs/07-customer/CLIENT_DATA_INTAKE.md):
- exact header contract; codes stay text so leading zeros survive; numbers become exact `Decimal`s and a value
  with more decimals than the SAP field allows is refused instead of rounded;
- SAP total rows (blank Plant, Material and Material type; one per unit of measure, one of them also carrying
  the value totals) are excluded from the data and kept as control totals — the detail rows must add up to them
  exactly or nothing is written;
- duplicates, stock without master data, zero prices and source values that differ from quantity × price are
  counted in the manifest and left for the data-quality checks, never silently fixed.

Error messages name rows and columns but never echo cell contents.
"""

from __future__ import annotations

import csv
import hashlib
import json
import os
import tempfile
from collections import Counter, defaultdict
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from enum import StrEnum
from pathlib import Path
from typing import Any

import duckdb
import yaml
from openpyxl import load_workbook

IMPORTER_VERSION = "sap-exports/1"
MAX_ERRORS = 20
NULL = r"\N"
DEFAULT_TOLERANCE = Decimal("0.01")
UNIT_PRICE_SCALE = 8
# SAP shows the price with 2 decimals per price unit while it stores the stock value itself, so quantity × price can
# differ from the stored value by up to |quantity| × half a cent ÷ price unit without any real difference.
PRICE_HALF_CENT = Decimal("0.005")


class IntakeError(Exception):
    def __init__(self, source: str, problems: list[str]) -> None:
        super().__init__(f"{source}: " + "; ".join(problems))
        self.source = source
        self.problems = problems


class Kind(StrEnum):
    CODE = "code"
    TEXT = "text"
    QUANTITY = "quantity"
    AMOUNT = "amount"
    INTEGER = "integer"
    DATE = "date"
    IGNORED = "ignored"  # read for the header contract only, never stored


@dataclass(frozen=True)
class Col:
    header: str
    name: str
    kind: Kind


# (scale, precision) — SAP quantity fields carry 3 decimals, currency amounts 2.
_NUMERIC = {Kind.QUANTITY: (3, 18), Kind.AMOUNT: (2, 18), Kind.INTEGER: (0, 9)}
SQL_TYPES = {
    Kind.CODE: "VARCHAR",
    Kind.TEXT: "VARCHAR",
    Kind.QUANTITY: "DECIMAL(18,3)",
    Kind.AMOUNT: "DECIMAL(18,2)",
    Kind.INTEGER: "INTEGER",
    Kind.DATE: "DATE",
}

C, T, Q, A, N, DT = Kind.CODE, Kind.TEXT, Kind.QUANTITY, Kind.AMOUNT, Kind.INTEGER, Kind.DATE

STOCK_COLUMNS: tuple[Col, ...] = (
    Col("Plant", "plant", C),
    Col("Storage Location", "storage_location", C),
    Col("Material type", "material_type", C),
    Col("Material Group", "material_group", C),
    Col("Material", "material", C),
    Col("Material Description", "material_description", T),
    Col("Base Unit of Measure", "base_uom", C),
    Col("Unrestricted", "unrestricted", Q),
    Col("Quality Inspection", "quality_inspection", Q),
    Col("Value in QualInsp.", "value_quality_inspection", A),
    Col("Transit and Transfer", "transit_transfer", Q),
    Col("Restricted-Use Stock", "restricted_use", Q),
    Col("Blocked", "blocked", Q),
    Col("Returns", "returns", Q),
    Col("Value Unrestricted", "value_unrestricted", A),
)
MASTER_COLUMNS: tuple[Col, ...] = (
    Col("Material", "material", C),
    Col("Plant", "plant", C),
    Col("Valuation Type", "valuation_type", C),
    Col("Material Description", "material_description", T),
    Col("Last Change", "last_change", DT),
    Col("Material type", "material_type", C),
    Col("Material Group", "material_group", C),
    Col("Base Unit of Measure", "base_uom", C),
    Col("Purchasing Group", "purchasing_group", C),
    Col("ABC Indicator", "abc_indicator", C),
    Col("MRP Type", "mrp_type", C),
    Col("Valuation Class", "valuation_class", C),
    Col("Price control", "price_control", C),
    Col("Price", "price", A),
    Col("Currency", "currency", C),
    Col("Price unit", "price_unit", N),
    Col("Created By", "created_by", Kind.IGNORED),  # a person's user id: not needed, not stored
)
# Stock status columns of the stock list, in file order, with the source value SAP reports for that status.
STATUSES: tuple[tuple[str, str, str | None], ...] = (
    ("unrestricted", "unrestricted", "value_unrestricted"),
    ("quality_inspection", "quality_inspection", "value_quality_inspection"),
    ("in_transit", "transit_transfer", None),
    ("restricted", "restricted_use", None),
    ("blocked", "blocked", None),
    ("returns", "returns", None),
)
QTY_COLUMNS = tuple(q for _, q, _ in STATUSES)
VALUE_COLUMNS = tuple(v for _, _, v in STATUSES if v)

Row = dict[str, Any]


# --------------------------------------------------------------------------------------- parsing
def _decimal(value: object, kind: Kind) -> Decimal:
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        raise ValueError("expected a number")
    try:
        d = Decimal(str(value))  # str() of a float is its shortest exact round-trip form
    except InvalidOperation as exc:
        raise ValueError("expected a number") from exc
    if not d.is_finite():
        raise ValueError("not a finite number")
    scale, precision = _NUMERIC[kind]
    q = d.quantize(Decimal(1).scaleb(-scale))
    if q != d:
        raise ValueError(f"more than {scale} decimal places (refused rather than rounded)")
    if abs(q) >= Decimal(10) ** (precision - scale):
        raise ValueError("number too large")
    return q + 0  # normalises -0


def parse_cell(value: object, kind: Kind) -> object:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    match kind:
        case Kind.CODE:
            if not isinstance(value, str):
                raise ValueError("a code arrived as a number, so leading zeros may be lost — export codes as text")
            return value.strip()
        case Kind.TEXT:
            if isinstance(value, bool) or not isinstance(value, (str, int)):
                raise ValueError("expected text")
            return str(value).strip()
        case Kind.QUANTITY | Kind.AMOUNT:
            return _decimal(value, kind)
        case Kind.INTEGER:
            return int(_decimal(value, kind))
        case Kind.DATE:
            if isinstance(value, datetime):
                return value.date()
            if isinstance(value, date):
                return value
            raise ValueError("expected a date")
    raise ValueError(f"unsupported column kind {kind}")


@dataclass
class Sheet:
    label: str
    path: Path
    sha256: str
    rows: list[tuple[int, Row]]  # (row number in the file, parsed values)
    blank_rows: int


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _blank(v: object) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def read_export(path: Path, columns: Sequence[Col], label: str) -> Sheet:
    """Read the first sheet of an export against its header contract."""
    if not path.is_file():
        raise IntakeError(label, [f"file not found: {path.name}"])
    digest = _sha256(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows_iter = workbook.worksheets[0].iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if header_row is None:
            raise IntakeError(label, ["the first sheet is empty"])
        header = ["" if h is None else str(h).strip() for h in header_row]
        while header and header[-1] == "":
            header.pop()
        expected = [c.header for c in columns]
        problems: list[str] = []
        if missing := [h for h in expected if h not in header]:
            problems.append("missing columns: " + ", ".join(missing))
        if extra := [h for h in header if h not in expected]:
            problems.append("unexpected columns: " + ", ".join(repr(h) for h in extra))
        if dups := sorted({h for h in header if header.count(h) > 1}):
            problems.append("duplicate columns: " + ", ".join(dups))
        if problems:
            raise IntakeError(label, [*problems, "expected exactly: " + ", ".join(expected)])
        index = {h: i for i, h in enumerate(header)}
        wanted = [c for c in columns if c.kind is not Kind.IGNORED]
        rows: list[tuple[int, Row]] = []
        errors: list[str] = []
        blank = 0
        for number, raw in enumerate(rows_iter, start=2):
            cells = list(raw)
            if all(_blank(v) for v in cells):
                blank += 1
                continue
            if any(not _blank(v) for v in cells[len(header) :]):
                errors.append(f"row {number}: values outside the header columns")
            cells += [None] * (len(header) - len(cells))
            parsed: Row = {}
            for col in wanted:
                try:
                    parsed[col.name] = parse_cell(cells[index[col.header]], col.kind)
                except ValueError as exc:
                    errors.append(f"row {number}, {col.header}: {exc}")
            rows.append((number, parsed))
            if len(errors) >= MAX_ERRORS:
                errors.append("… stopped after the first errors")
                break
        if errors:
            raise IntakeError(label, errors)
    finally:
        workbook.close()
    return Sheet(label, path, digest, rows, blank)


# --------------------------------------------------------------------------------------- checks
def split_totals(sheet: Sheet) -> tuple[list[tuple[int, Row]], list[tuple[int, Row]]]:
    """Separate SAP total rows from detail rows. Totals form one block after the last detail row."""
    details: list[tuple[int, Row]] = []
    totals: list[tuple[int, Row]] = []
    errors: list[str] = []
    for number, r in sheet.rows:
        if r["plant"] is None and r["material"] is None and r["material_type"] is None:
            filled = [k for k in ("storage_location", "material_group", "material_description") if r[k] is not None]
            if filled:
                errors.append(f"row {number}: looks like a total row but has {', '.join(filled)} filled")
            if r["base_uom"] is None:
                errors.append(f"row {number}: total row without a unit of measure")
            totals.append((number, r))
            continue
        if totals:
            errors.append(f"row {number}: detail row after the SAP total rows")
        if r["plant"] is None or r["material"] is None or r["base_uom"] is None:
            errors.append(f"row {number}: Plant, Material and Base Unit of Measure are required on detail rows")
        if any(r[q] is None for q in QTY_COLUMNS):
            errors.append(f"row {number}: a stock quantity is missing")
        details.append((number, r))
        if len(errors) >= MAX_ERRORS:
            break
    if errors:
        raise IntakeError(sheet.label, errors)
    return details, totals


def check_control_totals(details: list[tuple[int, Row]], totals: list[tuple[int, Row]]) -> dict[str, Any]:
    """The gate: every quantity column per unit and every value column must add up exactly to the SAP totals."""
    problems: list[str] = []
    if not totals:
        raise IntakeError(
            "stock", ["no SAP total rows: export the list with totals so the import can be checked against them"]
        )
    by_uom: dict[str, Row] = {}
    for number, r in totals:
        if r["base_uom"] in by_uom:
            problems.append(f"row {number}: second total row for one unit of measure")
        by_uom[r["base_uom"]] = r
    sums: dict[str, dict[str, Decimal]] = defaultdict(lambda: dict.fromkeys(QTY_COLUMNS, Decimal(0)))
    value_sums = dict.fromkeys(VALUE_COLUMNS, Decimal(0))
    for _, r in details:
        for q in QTY_COLUMNS:
            sums[r["base_uom"]][q] += r[q]
        for v in VALUE_COLUMNS:
            value_sums[v] += r[v] or 0
    if missing := sorted(set(sums) - set(by_uom), key=str):
        problems.append(f"{len(missing)} unit(s) of measure have detail rows but no total row")
    quantity_totals: dict[str, dict[str, str]] = {}
    for uom, t in sorted(by_uom.items(), key=lambda kv: str(kv[0])):
        got = sums.get(uom, dict.fromkeys(QTY_COLUMNS, Decimal(0)))
        for q in QTY_COLUMNS:
            if (t[q] or 0) != got[q]:
                problems.append(
                    f"{q} for one unit of measure: rows differ from the SAP total by {got[q] - (t[q] or 0)}"
                )
        quantity_totals[str(uom)] = {q: format(t[q] or Decimal(0), "f") for q in QTY_COLUMNS}
    carriers = [r for _, r in totals if any(r[v] is not None for v in VALUE_COLUMNS)]
    if len(carriers) != 1:
        problems.append(f"expected exactly one total row with value totals, found {len(carriers)}")
    value_totals: dict[str, str] = {}
    if len(carriers) == 1:
        for v in VALUE_COLUMNS:
            total = carriers[0][v]
            if total is None or total != value_sums[v]:
                diff = None if total is None else value_sums[v] - total
                problems.append(f"{v}: rows differ from the SAP total by {diff}")
            value_totals[v] = format(total or Decimal(0), "f")
    if problems:
        raise IntakeError("control totals (nothing was written)", problems)
    return {
        "passed": True,
        "rule": "Per unit of measure, each quantity column of the detail rows adds up exactly to the SAP total "
        "row; the value columns add up exactly to the grand total.",
        "units_of_measure": len(by_uom),
        "quantity_totals": quantity_totals,
        "value_totals": value_totals,
    }


def prepare_master(sheet: Sheet) -> tuple[list[tuple[int, Row]], dict[str, int]]:
    """Validate master keys and derive the unit price (Price ÷ Price unit) where it is exact and complete."""
    errors: list[str] = []
    keys: Counter[tuple[str, str]] = Counter()
    for number, r in sheet.rows:
        if r["material"] is None or r["plant"] is None:
            errors.append(f"row {number}: Material and Plant are required")
        keys[(r["material"], r["plant"])] += 1
    if dup := sum(1 for n in keys.values() if n > 1):
        errors.append(
            f"{dup} material/plant key(s) appear more than once (split valuation is not supported: one price "
            "per material and plant is needed to value stock)"
        )
    if errors:
        raise IntakeError(sheet.label, errors[:MAX_ERRORS])
    issues: Counter[str] = Counter()
    unit = Decimal(1).scaleb(-UNIT_PRICE_SCALE)
    for _, r in sheet.rows:
        price, per, currency = r["price"], r["price_unit"], r["currency"]
        r["unit_price"] = r["price_rounding"] = None
        if price is None:
            issues["no_price"] += 1
        elif per is None or per <= 0:
            issues["invalid_price_unit"] += 1
        elif currency is None:
            issues["price_without_currency"] += 1
        else:
            exact = (price / per).quantize(unit)
            if exact * per != price:
                issues["inexact_price_unit"] += 1
            else:
                r["unit_price"] = exact
                r["price_rounding"] = (PRICE_HALF_CENT / per).quantize(unit)
    return sheet.rows, dict(issues)


# --------------------------------------------------------------------------------------- writing
def _q(identifier: str) -> str:
    return '"' + identifier.replace('"', '""') + '"'


def _cell(value: object) -> str:
    if value is None:
        return NULL
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _load(
    con: duckdb.DuckDBPyConnection,
    table: str,
    columns: list[tuple[str, str]],
    rows: list[tuple[int, Row]],
    scratch: Path,
) -> None:
    ddl = ", ".join(f"{_q(name)} {sql_type}" for name, sql_type in columns)
    con.execute(f"CREATE TABLE {_q(table)} ({ddl})")
    path = scratch / f"{table}.csv"
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow([n for n, _ in columns])
        for number, r in rows:
            writer.writerow([_cell(number if n == "source_row" else r.get(n)) for n, _ in columns])
    con.execute(
        f"COPY {_q(table)} FROM ? (HEADER true, DELIMITER ',', QUOTE '\"', ESCAPE '\"', NULLSTR ?)",
        [str(path), NULL],
    )


def _columns(cols: Sequence[Col], *extra: tuple[str, str]) -> list[tuple[str, str]]:
    return [(c.name, SQL_TYPES[c.kind]) for c in cols if c.kind is not Kind.IGNORED] + list(extra)


_STOCK_BRANCHES = "\n    UNION ALL ".join(
    f"SELECT source_row, plant, storage_location, material, base_uom, '{status}' AS stock_status, "
    f"{qty} AS quantity, {value or 'CAST(NULL AS DECIMAL(18,2))'} AS source_value FROM sap_mb52"
    for status, qty, value in STATUSES
)
# Derived objects the tenant maps onto. The extract is an immutable snapshot, so they are materialised as tables at
# import time (a view would recompute the six-way stock-status unpivot and the master join on every query).
DERIVED = {
    # Long format: one line per stock status with a non-zero quantity or source value. The unit price is the
    # master Price ÷ Price unit for the same material and plant; stock without master data has no price.
    "v_stock": f"""
WITH s AS (
    {_STOCK_BRANCHES}
)
SELECT s.source_row, s.plant, s.storage_location, s.material, s.base_uom AS uom, s.stock_status, s.quantity,
       s.source_value, m.unit_price AS valuation_price, m.price_rounding, m.currency
FROM s LEFT JOIN sap_material_master m ON m.material = s.material AND m.plant = s.plant
WHERE s.quantity <> 0 OR coalesce(s.source_value, 0) <> 0""",
    # Client-level attributes: from the master where present, otherwise from the stock list.
    "v_item": """
WITH m AS (
    SELECT material, min(material_description) AS description, min(material_type) AS material_type,
           min(material_group) AS material_group, min(base_uom) AS base_uom
    FROM sap_material_master GROUP BY material
), s AS (
    SELECT material, min(material_description) AS description, min(material_type) AS material_type,
           min(material_group) AS material_group, min(base_uom) AS base_uom
    FROM sap_mb52 GROUP BY material
)
SELECT coalesce(m.material, s.material) AS material, coalesce(m.description, s.description) AS description,
       coalesce(m.material_type, s.material_type) AS material_type,
       coalesce(m.material_group, s.material_group) AS material_group, coalesce(m.base_uom, s.base_uom) AS base_uom
FROM m FULL OUTER JOIN s ON s.material = m.material""",
    "v_item_warehouse": """
SELECT material, plant, abc_indicator, mrp_type, price_control, valuation_class, purchasing_group,
       unit_price AS valuation_price, currency
FROM sap_material_master""",
    "v_warehouse": "SELECT plant FROM sap_mb52 UNION SELECT plant FROM sap_material_master",
}
DATASETS = ("item", "item_warehouse", "stock_balance", "warehouse")


def _scalar(con: duckdb.DuckDBPyConnection, sql: str, params: list[Any] | None = None) -> Any:
    row = con.execute(sql, params or []).fetchone()
    return None if row is None else row[0]


def _profile(con: duckdb.DuckDBPyConnection, tolerance: Decimal) -> dict[str, Any]:
    """Counts only (no amounts) describing what the checks will find."""

    def n(sql: str, params: list[Any] | None = None) -> int:
        return int(_scalar(con, sql, params) or 0)

    stocked = " OR ".join(f"s.{q} <> 0" for q in QTY_COLUMNS)
    joined = "sap_mb52 s LEFT JOIN sap_material_master m ON m.material = s.material AND m.plant = s.plant"
    lines = dict(con.execute("SELECT stock_status, count(*) FROM v_stock GROUP BY 1 ORDER BY 1").fetchall())
    valued = "source_value IS NOT NULL AND valuation_price IS NOT NULL"
    diff = "abs(source_value - quantity * valuation_price)"
    bound = "? + abs(quantity) * price_rounding"
    return {
        "stock": {
            "detail_rows": n("SELECT count(*) FROM sap_mb52"),
            "all_zero_rows": n(
                f"SELECT count(*) FROM sap_mb52 s WHERE NOT ({stocked}) "
                "AND coalesce(value_unrestricted, 0) = 0 AND coalesce(value_quality_inspection, 0) = 0"
            ),
            "plants": n("SELECT count(DISTINCT plant) FROM sap_mb52"),
            "storage_locations": n("SELECT count(DISTINCT (plant, storage_location)) FROM sap_mb52"),
            "materials": n("SELECT count(DISTINCT material) FROM sap_mb52"),
            "units_of_measure": n("SELECT count(DISTINCT base_uom) FROM sap_mb52"),
            "duplicate_keys": n(
                "SELECT count(*) FROM (SELECT 1 FROM sap_mb52 GROUP BY plant, storage_location, material "
                "HAVING count(*) > 1)"
            ),
            "rows_without_master": n(f"SELECT count(*) FROM {joined} WHERE m.material IS NULL"),
            "stocked_rows_without_master": n(f"SELECT count(*) FROM {joined} WHERE m.material IS NULL AND ({stocked})"),
            "stocked_rows_zero_price": n(f"SELECT count(*) FROM {joined} WHERE m.price = 0 AND ({stocked})"),
            "stocked_rows_without_unit_price": n(
                f"SELECT count(*) FROM {joined} WHERE m.material IS NOT NULL AND m.unit_price IS NULL AND ({stocked})"
            ),
            "uom_differs_from_master": n(f"SELECT count(*) FROM {joined} WHERE m.base_uom <> s.base_uom"),
            "stock_lines_by_status": {k: int(v) for k, v in lines.items()},
        },
        "master": {
            "rows": n("SELECT count(*) FROM sap_material_master"),
            "plants": n("SELECT count(DISTINCT plant) FROM sap_material_master"),
            "materials": n("SELECT count(DISTINCT material) FROM sap_material_master"),
            "currencies": n("SELECT count(DISTINCT currency) FROM sap_material_master"),
            "zero_price": n("SELECT count(*) FROM sap_material_master WHERE price = 0"),
            "abc_blank": n("SELECT count(*) FROM sap_material_master WHERE abc_indicator IS NULL"),
            "mrp_type_nd": n("SELECT count(*) FROM sap_material_master WHERE mrp_type = 'ND'"),
            "price_control": {
                str(k): int(v)
                for k, v in con.execute(
                    "SELECT coalesce(price_control, '(blank)'), count(*) FROM sap_material_master GROUP BY 1 ORDER BY 1"
                ).fetchall()
            },
            "price_units": {
                str(k): int(v)
                for k, v in con.execute(
                    "SELECT price_unit, count(*) FROM sap_material_master GROUP BY 1 ORDER BY 1"
                ).fetchall()
            },
            "records_without_stock": n(
                "SELECT count(*) FROM sap_material_master m WHERE NOT EXISTS (SELECT 1 FROM sap_mb52 s "
                f"WHERE s.material = m.material AND s.plant = m.plant AND ({stocked}))"
            ),
            "materials_with_inconsistent_attributes": n(
                "SELECT count(*) FROM (SELECT material FROM sap_material_master GROUP BY material HAVING "
                "count(DISTINCT material_type) > 1 OR count(DISTINCT base_uom) > 1 "
                "OR count(DISTINCT material_group) > 1)"
            ),
        },
        "reconciliation": {
            "rule": "Per stock line with a source value and a unit price: |source value − quantity × (Price ÷ Price "
            "unit)| ≤ tolerance + |quantity| × 0.005 ÷ Price unit (the price is shown rounded to 2 decimals).",
            "tolerance": format(tolerance, "f"),
            "compared_lines": n(f"SELECT count(*) FROM v_stock WHERE {valued}"),
            "matching_lines": n(f"SELECT count(*) FROM v_stock WHERE {valued} AND {diff} <= {bound}", [tolerance]),
            "differing_lines": n(f"SELECT count(*) FROM v_stock WHERE {valued} AND {diff} > {bound}", [tolerance]),
            "within_price_rounding_only": n(
                f"SELECT count(*) FROM v_stock WHERE {valued} AND {diff} > ? AND {diff} <= {bound}",
                [tolerance, tolerance],
            ),
            "valued_lines_without_price": n(
                "SELECT count(*) FROM v_stock WHERE coalesce(source_value, 0) <> 0 AND valuation_price IS NULL"
            ),
        },
    }


@dataclass(frozen=True)
class ImportResult:
    target: Path
    manifest: dict[str, Any]
    plants: list[str]


def import_sap_exports(
    stock_path: Path,
    master_path: Path,
    *,
    tenant_id: str,
    as_of: date,
    target: Path,
    synthetic: bool,
    manifest_path: Path | None = None,
    tolerance: Decimal = DEFAULT_TOLERANCE,
) -> ImportResult:
    """Validate both files, run the control-total gate, then write `target` atomically (stop the API first)."""
    stock = read_export(stock_path, STOCK_COLUMNS, "stock")
    master = read_export(master_path, MASTER_COLUMNS, "master")
    details, totals = split_totals(stock)
    control = check_control_totals(details, totals)
    master_rows, valuation_issues = prepare_master(master)

    target.parent.mkdir(parents=True, exist_ok=True)
    tmp_db = target.with_suffix(".tmp.duckdb")
    for leftover in (tmp_db, tmp_db.with_suffix(".duckdb.wal")):
        leftover.unlink(missing_ok=True)
    con = duckdb.connect(str(tmp_db))
    try:
        with tempfile.TemporaryDirectory(prefix="lumiant-ingest-") as scratch:
            row_col = ("source_row", "INTEGER")
            _load(con, "sap_mb52", _columns(STOCK_COLUMNS, row_col), details, Path(scratch))
            _load(con, "sap_mb52_totals", _columns(STOCK_COLUMNS, row_col), totals, Path(scratch))
            _load(
                con,
                "sap_material_master",
                _columns(
                    MASTER_COLUMNS,
                    ("unit_price", f"DECIMAL(20,{UNIT_PRICE_SCALE})"),
                    ("price_rounding", f"DECIMAL(20,{UNIT_PRICE_SCALE})"),
                    row_col,
                ),
                master_rows,
                Path(scratch),
            )
        for name, sql in DERIVED.items():
            con.execute(f"CREATE TABLE {_q(name)} AS {sql}")
        plants = [str(p) for (p,) in con.execute("SELECT plant FROM v_warehouse ORDER BY plant").fetchall()]
        manifest: dict[str, Any] = {
            "importer_version": IMPORTER_VERSION,
            "tenant_id": tenant_id,
            "synthetic": synthetic,
            "as_of": as_of.isoformat(),
            "imported_at": datetime.now(UTC).isoformat(timespec="seconds"),
            "inputs": {
                s.label: {"file": s.path.name, "sha256": s.sha256, "rows": len(s.rows), "blank_rows": s.blank_rows}
                for s in (stock, master)
            },
            "stock_total_rows_excluded": len(totals),
            "control_totals": control,
            "valuation_issues": valuation_issues,
            **_profile(con, tolerance),
        }
        con.execute("CREATE TABLE _dataset_info (key VARCHAR, value VARCHAR)")
        info = {
            "synthetic": "true" if synthetic else "false",
            "tenant_id": tenant_id,
            "as_of": as_of.isoformat(),
            "importer_version": IMPORTER_VERSION,
            "manifest": json.dumps(manifest, sort_keys=True),
        }
        con.executemany("INSERT INTO _dataset_info VALUES (?, ?)", list(info.items()))
        con.execute("CREATE TABLE _extract_metadata (dataset VARCHAR, extracted_at TIMESTAMP)")
        extracted = datetime(as_of.year, as_of.month, as_of.day)
        con.executemany("INSERT INTO _extract_metadata VALUES (?, ?)", [(d, extracted) for d in DATASETS])
        con.execute("CHECKPOINT")
    finally:
        con.close()
    try:
        os.replace(tmp_db, target)
    except PermissionError as exc:
        tmp_db.unlink(missing_ok=True)
        raise IntakeError(str(target), ["file is in use — stop the API before importing"]) from exc
    if manifest_path is not None:
        _atomic_write(manifest_path, json.dumps(manifest, indent=2, sort_keys=True) + "\n")
    return ImportResult(target, manifest, plants)


def _atomic_write(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(text, encoding="utf-8")
    os.replace(tmp, path)


# --------------------------------------------------------------------------------------- tenant config
# SAP *standard* meanings only (customising may differ per client — docs/07-customer/OPEN_QUESTIONS.md Q-019,
# Q-021). Codes not listed read as "unmapped" and are reported by the data-quality checks.
SAP_MATERIAL_TYPES = {
    "ROH": "RAW",
    "HALB": "SEMI",
    "FERT": "FINISHED",
    "HAWA": "TRADING",
    "ERSA": "SPARE",
    "VERP": "PACKAGING",
    "HIBE": "CONSUMABLE",
    "NLAG": "NON_STOCK",
    "UNBW": "NON_VALUATED",
    "DIEN": "SERVICE",
    "FHMI": "TOOLING",
}
SAP_MRP_TYPES = {
    "ND": "none",
    "VB": "reorder_point",
    "VM": "reorder_point",
    "V1": "reorder_point",
    "V2": "reorder_point",
    "VV": "forecast",
    "PD": "mrp",
    "M0": "mrp",
    "M1": "mrp",
    "M2": "mrp",
    "M3": "mrp",
    "M4": "mrp",
}
SAP_PRICE_CONTROL = {"V": "moving_average", "S": "standard"}
SAP_SOURCE_HINTS = {
    "inventory_transaction": "SAP MB51 / MSEG material movements for the last 24 months",
    "reorder_policy": "SAP MARC plant planning data (reorder point, safety stock, maximum stock)",
    "lot": "SAP MCHB / MCH1 batch stock with shelf-life expiration dates",
    "stock_balance.lot_id": "SAP MCHB batch stock (stock per batch)",
    "purchase_record": "SAP EKKO / EKPO purchase order lines (open and received)",
    "supplier": "SAP LFA1 vendor master",
    "warehouse.name": "SAP T001W plant names",
    "storage_location": "SAP T001L storage location names",
    "item.abc_class": "ABC class per plant is in the material master (item_warehouse.abc_class)",
}
ROLES = {
    "super_admin": [
        "dashboard:read",
        "inventory:read",
        "cost:read",
        "query:run",
        "query:view_sql",
        "audit:read",
        "admin:tenant",
        "chat:use",
        "export:data",
    ],
    "standard_user": ["dashboard:read", "inventory:read", "cost:read", "query:run", "chat:use", "export:data"],
    "viewer_no_cost": ["dashboard:read", "inventory:read", "chat:use"],
}
DEPARTMENT = {"id": "PLANTS", "name": "All plants (organisation structure not provided in the extract)"}


def sap_tenant_config(
    *,
    tenant_id: str,
    display_name: str,
    data_classification: str,
    source_system: str,
    connection_path: str,
    plants: list[str],
    dev_users: list[dict[str, Any]],
) -> dict[str, Any]:
    """The tenant mapping for databases written by `import_sap_exports` (same shape for client and twin)."""
    return {
        "id": tenant_id,
        "display_name": display_name,
        "data_classification": data_classification,
        "source_system": source_system,
        "connection": {"type": "duckdb", "path": connection_path},
        "freshness": {"object": "_extract_metadata", "column": "extracted_at"},
        "policies": {"max_rows": 1000, "default_rows": 100, "query_timeout_s": 15, "stale_after_hours": 36},
        "ai": {"external_provider_approved": False},
        "departments": [DEPARTMENT],
        "warehouses": [{"id": p, "department": DEPARTMENT["id"]} for p in plants],
        "roles": ROLES,
        "dev_users": dev_users,
        "datasets": {
            "item": {
                "source": "v_item",
                "fields": {
                    "item_id": "material",
                    "item_code": "material",
                    "description": "description",
                    "item_group": "material_group",
                    "inventory_stage": "material_type",
                    "base_uom": "base_uom",
                },
                "attributes": {
                    "material_type": {"column": "material_type", "type": "string", "description": "SAP material type"}
                },
                "value_maps": {"inventory_stage": SAP_MATERIAL_TYPES},
            },
            "warehouse": {"source": "v_warehouse", "fields": {"warehouse_id": "plant"}},
            "item_warehouse": {
                "source": "v_item_warehouse",
                "fields": {
                    "item_id": "material",
                    "warehouse_id": "plant",
                    "abc_class": "abc_indicator",
                    "planning_type": "mrp_type",
                    "price_control": "price_control",
                    "valuation_price": "valuation_price",
                    "currency": "currency",
                },
                "attributes": {
                    "mrp_type": {"column": "mrp_type", "type": "string", "description": "SAP MRP type"},
                    "valuation_class": {
                        "column": "valuation_class",
                        "type": "string",
                        "description": "SAP valuation class",
                    },
                    "purchasing_group": {
                        "column": "purchasing_group",
                        "type": "string",
                        "description": "SAP purchasing group",
                    },
                },
                "value_maps": {"planning_type": SAP_MRP_TYPES, "price_control": SAP_PRICE_CONTROL},
            },
            "stock_balance": {
                "source": "v_stock",
                "fields": {
                    "item_id": "material",
                    "warehouse_id": "plant",
                    "location_id": "storage_location",
                    "stock_status": "stock_status",
                    "quantity": "quantity",
                    "uom": "uom",
                    "valuation_price": "valuation_price",
                    "currency": "currency",
                    "source_value": "source_value",
                    "price_rounding": "price_rounding",
                },
            },
        },
        "source_hints": SAP_SOURCE_HINTS,
    }


def dump_tenant_yaml(config: dict[str, Any], header: str) -> str:
    comment = "".join(f"# {line}\n" if line else "#\n" for line in header.splitlines())
    return comment + yaml.safe_dump(config, sort_keys=False, allow_unicode=True, width=120)


def write_tenant_yaml(path: Path, config: dict[str, Any], header: str) -> None:
    _atomic_write(path, dump_tenant_yaml(config, header))
