"""Reports: the same secured analyses, shaped as downloadable/printable tables with a metadata header.

Exports are authorised by the backend (`export:data`), audited, watermarked when the data is synthetic or a
confidential client extract, and protected against spreadsheet formula injection."""

from __future__ import annotations

import csv
import io
import re
from collections.abc import Callable
from dataclasses import dataclass
from datetime import UTC, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from lumiant.analyses.below_reorder import BelowReorderFilters, run_below_reorder
from lumiant.analyses.risks import run_dormant, run_expiring
from lumiant.auth.permissions import Permission
from lumiant.auth.principal import Principal
from lumiant.core.errors import ForbiddenError, NotFoundError
from lumiant.query.plan import QueryPlan
from lumiant.query.results import QueryResult
from lumiant.query.service import QueryService

SYNTHETIC_NOTICE = "SYNTHETIC DEMO DATA - generated for demonstration, not from any customer system"
DRAFT_NOTICE = "Calculation definitions are drafts until approved by the customer"
CONFIDENTIAL_NOTICE = "CONFIDENTIAL - client data extract for local analysis only; do not distribute (ADR-0008)"
_NUMERIC = re.compile(r"^-?\d+(\.\d+)?$")


@dataclass(frozen=True)
class ReportDef:
    id: str
    title: str
    description: str
    permissions: tuple[Permission, ...]
    columns: tuple[tuple[str, str], ...]  # (result key, header)
    build: Callable[[QueryService, Principal, dict[str, int], int], QueryResult]
    params: tuple[str, ...] = ()
    draft: bool = False


def _plan(
    query: QueryService, principal: Principal, plan: dict[str, Any], action: str, optional: tuple[str, ...] = ()
) -> QueryResult:
    """`optional` dimensions are descriptive columns dropped when the tenant's source does not provide them."""
    mapping = query.mapping(principal.tenant_id)
    dims = [d for d in plan.get("dimensions", []) if d not in optional or mapping.has_ref(plan["dataset"], d)]
    return query.run_plan(QueryPlan.model_validate({**plan, "dimensions": dims}), principal, action=action)


def _value_by_warehouse(q: QueryService, p: Principal, _: dict[str, int], limit: int) -> QueryResult:
    return _plan(
        q,
        p,
        {
            "dataset": "stock_balance",
            "dimensions": ["warehouse_id", "warehouse.name"],
            "measures": [{"field": "stock_value", "agg": "sum", "alias": "value"}],
            "order_by": [{"field": "value", "direction": "desc"}],
            "limit": limit,
        },
        "report:value_by_warehouse",
        ("warehouse.name",),
    )


def _stock_listing(q: QueryService, p: Principal, _: dict[str, int], limit: int) -> QueryResult:
    dims = [
        "item.item_code",
        "item.description",
        "item.item_group",
        "warehouse_id",
        "location_id",
        "lot_id",
        "stock_status",
        "quantity",
    ]
    if p.has(Permission.COST_READ):
        dims += ["valuation_price", "stock_value"]
    return _plan(
        q,
        p,
        {"dataset": "stock_balance", "dimensions": dims, "order_by": [{"field": "item.item_code"}], "limit": limit},
        "report:stock_listing",
        (
            "item.description",
            "item.item_group",
            "location_id",
            "lot_id",
            "stock_status",
            "valuation_price",
            "stock_value",
        ),
    )


def _below(q: QueryService, p: Principal, _: dict[str, int], limit: int) -> QueryResult:
    return run_below_reorder(q, p, BelowReorderFilters(), limit=limit, offset=0)


def _excess(q: QueryService, p: Principal, _: dict[str, int], limit: int) -> QueryResult:
    return run_below_reorder(q, p, BelowReorderFilters(), limit=limit, offset=0, mode="excess")


def _dormant(q: QueryService, p: Principal, params: dict[str, int], limit: int) -> QueryResult:
    return run_dormant(q, p, days=params.get("days", 270), limit=limit, offset=0)


def _expiring(q: QueryService, p: Principal, params: dict[str, int], limit: int) -> QueryResult:
    return run_expiring(q, p, days=params.get("days", 60), limit=limit)


INV, COST, EXPORT = Permission.INVENTORY_READ, Permission.COST_READ, Permission.EXPORT_DATA
_ITEM = (("item_code", "Item"), ("description", "Description"), ("warehouse_id", "Warehouse"))

REPORTS: dict[str, ReportDef] = {
    r.id: r
    for r in (
        ReportDef(
            "value_by_warehouse",
            "Inventory value by warehouse",
            "On-hand value per warehouse and currency.",
            (INV, COST),
            (("warehouse_id", "Warehouse"), ("warehouse.name", "Name"), ("currency", "Currency"), ("value", "Value")),
            _value_by_warehouse,
        ),
        ReportDef(
            "stock_listing",
            "Stock listing",
            "Every stock line in your scope.",
            (INV,),
            (
                ("item.item_code", "Item"),
                ("item.description", "Description"),
                ("item.item_group", "Group"),
                ("warehouse_id", "Warehouse"),
                ("location_id", "Location"),
                ("lot_id", "Lot"),
                ("stock_status", "Stock status"),
                ("quantity", "Quantity"),
                ("uom", "Unit"),
                ("currency", "Currency"),
                ("valuation_price", "Valuation price"),
                ("stock_value", "Value"),
            ),
            _stock_listing,
        ),
        ReportDef(
            "below_reorder",
            "Items below reorder level",
            "Customer reorder policy versus on-hand.",
            (INV,),
            (
                *_ITEM,
                ("uom", "Unit"),
                ("on_hand", "On hand"),
                ("reorder_level", "Reorder level"),
                ("shortfall", "Shortfall"),
            ),
            _below,
        ),
        ReportDef(
            "excess_stock",
            "Excess stock",
            "On-hand above the customer maximum.",
            (INV,),
            (*_ITEM, ("uom", "Unit"), ("on_hand", "On hand"), ("max_stock", "Maximum"), ("excess", "Excess")),
            _excess,
        ),
        ReportDef(
            "dormant_stock",
            "Dormant stock",
            "Stock with no issue in the chosen window.",
            (INV,),
            (
                *_ITEM,
                ("uom", "Unit"),
                ("on_hand", "On hand"),
                ("last_issue", "Last issue"),
                ("days_since_issue", "Days idle"),
            ),
            _dormant,
            ("days",),
        ),
        ReportDef(
            "expiring_lots",
            "Expiring lots",
            "Lots expired or expiring in the chosen window.",
            (INV,),
            (
                ("item.item_code", "Item"),
                ("warehouse_id", "Warehouse"),
                ("lot_id", "Lot"),
                ("quantity", "Quantity"),
                ("uom", "Unit"),
                ("lot.expiry_date", "Expires"),
                ("days_to_expiry", "Days left"),
            ),
            _expiring,
            ("days",),
        ),
    )
}


def safe_cell(value: Any) -> Any:
    """Neutralise spreadsheet formula injection: text that starts like a formula is prefixed with an apostrophe."""
    if isinstance(value, str) and value and value[0] in "=+-@\t\r" and not _NUMERIC.match(value):
        return "'" + value
    return value


class ReportService:
    def __init__(self, query: QueryService) -> None:
        self.query = query

    def visible(self, principal: Principal) -> list[ReportDef]:
        return [r for r in REPORTS.values() if all(principal.has(p) for p in r.permissions)]

    def _get(self, report_id: str, principal: Principal) -> ReportDef:
        report = REPORTS.get(report_id)
        if report is None:
            raise NotFoundError("report_not_found", "Unknown report.")
        if not all(principal.has(p) for p in report.permissions):
            self.query.record(principal, f"report:{report_id}", None, "deny", "denied", reason="missing_permission")
            raise ForbiddenError("forbidden", "Your role cannot view this report.")
        return report

    def build(
        self, report_id: str, principal: Principal, params: dict[str, int], *, limit: int | None = None
    ) -> dict[str, Any]:
        report = self._get(report_id, principal)
        max_rows = self.query.tenants.get(principal.tenant_id).policies.max_rows
        result = report.build(self.query, principal, params, min(limit or max_rows, max_rows))
        payload = result.payload()
        rows = payload["data"]
        keys = [(k, h) for k, h in report.columns if any(k in r for r in rows) or not rows]
        meta = self.query.tenant_meta(principal.tenant_id)
        prov = payload["provenance"]
        filters = [f"{f.get('field')} {f.get('op')} {f.get('value')}" for f in prov["filters"]["user"] if "field" in f]
        injected = prov["filters"]["injected"]
        metadata = {
            "title": report.title,
            "tenant": meta.display_name,
            "source_system": meta.source_system,
            "generated_at": datetime.now(UTC).isoformat(timespec="seconds"),
            "data_as_of": meta.data_as_of.isoformat() if meta.data_as_of else "unknown",
            "user": principal.user_id,
            "parameters": params,
            "filters": filters or ["none"],
            "access_scope": f"{len(injected[0]['value'])} warehouses (permission scope)"
            if injected
            else "all warehouses in tenant",
            "rows": len(rows),
            "truncated": bool(payload["provenance"].get("truncated")),
            "synthetic": meta.synthetic,
            "data_classification": meta.data_classification,
            "notices": [
                n
                for n in (
                    CONFIDENTIAL_NOTICE if meta.data_classification == "customer_extract" else None,
                    SYNTHETIC_NOTICE if meta.synthetic else None,
                    DRAFT_NOTICE
                    if (
                        report.draft
                        or report.id
                        in ("below_reorder", "excess_stock", "dormant_stock", "expiring_lots", "value_by_warehouse")
                    )
                    else None,
                )
                if n
            ],
            "currency_note": "Amounts in different currencies are listed separately and never added together.",
            "source_objects": prov["source_objects"],
        }
        return {
            "report": {"id": report.id, "title": report.title, "params": list(report.params)},
            "columns": [{"key": k, "label": h} for k, h in keys],
            "rows": rows,
            "metadata": metadata,
            "warnings": payload["warnings"],
            "provenance": prov,
        }

    def export(self, report_id: str, principal: Principal, params: dict[str, int], fmt: str) -> tuple[bytes, str, str]:
        if not principal.has(Permission.EXPORT_DATA):
            self.query.record(principal, f"export:{report_id}", None, "deny", "denied", reason="missing_permission")
            raise ForbiddenError("forbidden", "Your role cannot export data.")
        built = self.build(report_id, principal, params)
        meta, cols, rows = built["metadata"], built["columns"], built["rows"]
        header_lines = [
            ("Report", meta["title"]),
            ("Organisation", meta["tenant"]),
            ("Source", meta["source_system"]),
            ("Generated at (UTC)", meta["generated_at"]),
            ("Data as of", meta["data_as_of"]),
            ("Generated by", meta["user"]),
            ("Parameters", ", ".join(f"{k}={v}" for k, v in meta["parameters"].items()) or "none"),
            ("Filters", "; ".join(meta["filters"])),
            ("Access scope", meta["access_scope"]),
            ("Rows", f"{meta['rows']}{' (TRUNCATED - more rows exist)' if meta['truncated'] else ''}"),
            ("Currency", meta["currency_note"]),
            *[("NOTICE", n) for n in meta["notices"]],
        ]
        stamp = datetime.now(UTC).strftime("%Y%m%d-%H%M")
        prefix = "CONFIDENTIAL-" if meta["data_classification"] == "customer_extract" else ""
        name = f"{prefix}{report_id}-{stamp}.{fmt}"
        self.query.record(principal, f"export:{report_id}", None, "allow", "ok", rows=len(rows), reason=fmt)
        if fmt == "csv":
            buf = io.StringIO()
            w = csv.writer(buf)
            for k, v in header_lines:
                w.writerow([f"# {k}", safe_cell(str(v))])
            w.writerow([])
            w.writerow([c["label"] for c in cols])
            for r in rows:
                w.writerow([safe_cell(r.get(c["key"], "")) if r.get(c["key"]) is not None else "" for c in cols])
            return ("﻿" + buf.getvalue()).encode("utf-8"), "text/csv; charset=utf-8", name
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "Report"
        for k, v in header_lines:
            ws.append([k, safe_cell(str(v))])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        ws.append([])
        ws.append([c["label"] for c in cols])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
        for r in rows:
            # Decimals stay text: spreadsheets would convert them to floating point.
            ws.append([safe_cell(r.get(c["key"])) if r.get(c["key"]) is not None else None for c in cols])
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", name
