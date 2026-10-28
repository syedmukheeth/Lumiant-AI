"""SAP-extract pipeline end to end (M8): the demo-gamma twin goes through the same importer as the client files.

Expected numbers are computed straight from the generated XLSX files with openpyxl and Decimal, independently of
the importer's tables and views. The second half checks the customer_extract guards (ADR-0008) over HTTP.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from collections.abc import Iterator
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from lumiant.analyses.risks import run_data_quality
from lumiant.auth.principal import Principal
from lumiant.core.config import Settings
from lumiant.ingest.sap_exports import SAP_MATERIAL_TYPES, import_sap_exports, sap_tenant_config, write_tenant_yaml
from lumiant.main import create_app
from lumiant.query.plan import QueryPlan
from lumiant.query.service import QueryService
from lumiant.synthetic.sap_twin import SAP_TWINS, write_twin_exports
from lumiant.tenancy.registry import TenantRegistry
from tests.conftest import AS_OF, TEST_ITEMS
from tests.helpers import app_db_url

pytestmark = pytest.mark.integration
CSRF = {"X-Requested-With": "lumiant"}
STATUS_COLUMNS = {
    "unrestricted": ("Unrestricted", "Value Unrestricted"),
    "quality_inspection": ("Quality Inspection", "Value in QualInsp."),
    "in_transit": ("Transit and Transfer", None),
    "restricted": ("Restricted-Use Stock", None),
    "blocked": ("Blocked", None),
    "returns": ("Returns", None),
}
TOL = Decimal("0.01")
HALF_CENT = Decimal("0.005")


def dec(v: Any) -> Decimal:
    return Decimal(0) if v is None or v == "" else Decimal(str(v))


def read(path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    wb.close()
    width = len(rows[0])  # read-only mode drops trailing empty cells
    return [dict(zip(rows[0], (*r, *(None,) * (width - len(r))), strict=True)) for r in rows[1:]]


@pytest.fixture(scope="module")
def files(tmp_path_factory: pytest.TempPathFactory) -> dict[str, Any]:
    """The same deterministic twin files the session fixture imported, parsed independently."""
    stock, master, _ = write_twin_exports(
        tmp_path_factory.mktemp("twin"), SAP_TWINS["demo-gamma"], as_of=AS_OF, n_materials=TEST_ITEMS["demo-gamma"]
    )
    rows = read(stock)
    details = [r for r in rows if r["Plant"]]
    footer = [r for r in rows if not r["Plant"]]
    prices: dict[tuple[str, str], dict[str, Any]] = {}
    for m in read(master):
        unit = None
        if m["Price"] is not None and m["Price unit"]:
            u = dec(m["Price"]) / Decimal(m["Price unit"])
            unit = u if u.quantize(Decimal("1e-8")) * m["Price unit"] == dec(m["Price"]) else None  # else unpriced
        prices[(m["Material"], m["Plant"])] = {**m, "unit": unit}
    lines = []  # one per stock row and status with a non-zero quantity or source value, as SAP lists it
    for r in details:
        for status, (qcol, vcol) in STATUS_COLUMNS.items():
            q, v = dec(r[qcol]), (dec(r[vcol]) if vcol else None)
            if q != 0 or (v or 0) != 0:
                lines.append({"row": r, "status": status, "q": q, "v": v, "m": prices.get((r["Material"], r["Plant"]))})
    return {"details": details, "footer": footer, "master": list(prices.values()), "lines": lines}


def run(svc: QueryService, who: Principal, plan: dict[str, Any]) -> list[dict[str, Any]]:
    return svc.run_plan(QueryPlan.model_validate(plan), who, action="test").rows


# --------------------------------------------------------------------------------------- totals
def test_gamma_value_matches_the_files(
    query_service: QueryService, principals: dict[str, Principal], files: dict[str, Any]
) -> None:
    expected: dict[str, Decimal] = defaultdict(Decimal)
    for ln in files["lines"]:
        if ln["m"] and ln["m"]["unit"] is not None:
            expected[ln["m"]["Currency"]] += ln["q"] * ln["m"]["unit"]
    rows = run(
        query_service,
        principals["gamma.admin"],
        {"dataset": "stock_balance", "measures": [{"field": "stock_value", "agg": "sum", "alias": "value"}]},
    )
    got = {r["currency"]: r["value"] for r in rows if r["value"] is not None}
    assert got == dict(expected)
    assert set(got) == {"SAR"}


def test_gamma_source_values_equal_the_sap_control_total(
    query_service: QueryService, principals: dict[str, Principal], files: dict[str, Any]
) -> None:
    carrier = next(f for f in files["footer"] if f["Value Unrestricted"] is not None)
    sap_total = dec(carrier["Value Unrestricted"]) + dec(carrier["Value in QualInsp."])
    rows = run(
        query_service,
        principals["gamma.admin"],
        {"dataset": "stock_balance", "measures": [{"field": "source_value", "agg": "sum", "alias": "v"}]},
    )
    assert sum((r["v"] for r in rows if r["v"] is not None), Decimal(0)) == sap_total


def test_gamma_quantities_per_unit_and_status_equal_the_sap_totals(
    query_service: QueryService, principals: dict[str, Principal], files: dict[str, Any]
) -> None:
    rows = run(
        query_service,
        principals["gamma.admin"],
        {
            "dataset": "stock_balance",
            "dimensions": ["stock_status"],
            "measures": [{"field": "quantity", "agg": "sum", "alias": "q"}],
            "limit": 500,
        },
    )
    got = {(r["uom"], r["stock_status"]): r["q"] for r in rows}
    expected = {
        (f["Base Unit of Measure"], status): dec(f[qcol])
        for f in files["footer"]
        for status, (qcol, _) in STATUS_COLUMNS.items()
        if dec(f[qcol]) != 0
    }
    assert {k: v for k, v in got.items() if v != 0} == expected


# --------------------------------------------------------------------------------------- data quality
def test_gamma_data_quality_counts_match_the_files(
    query_service: QueryService, principals: dict[str, Principal], files: dict[str, Any]
) -> None:
    lines, master = files["lines"], files["master"]
    stocked = {(ln["row"]["Material"], ln["row"]["Plant"]) for ln in lines if ln["q"] != 0}
    keys = Counter((ln["row"]["Material"], ln["row"]["Storage Location"], ln["status"]) for ln in lines)
    types = {r["Material"]: r["Material type"] for r in files["details"]} | {
        m["Material"]: m["Material type"] for m in master
    }
    expected = {
        "stock_without_master": sum(1 for ln in lines if ln["q"] != 0 and ln["m"] is None),
        "zero_valuation_price": sum(1 for ln in lines if ln["q"] != 0 and ln["m"] and ln["m"]["unit"] == 0),
        "source_value_mismatch": sum(
            1
            for ln in lines
            if ln["v"] is not None
            and ln["m"]
            and ln["m"]["unit"] is not None
            # SAP shows the price rounded to 2 decimals per price unit: that much difference is not a mismatch.
            and abs(ln["v"] - ln["q"] * ln["m"]["unit"]) > TOL + abs(ln["q"]) * HALF_CENT / ln["m"]["Price unit"]
        ),
        "duplicate_stock_keys": sum(1 for n in keys.values() if n > 1),
        "unmapped_item_inventory_stage": sum(1 for t in types.values() if t and t not in SAP_MATERIAL_TYPES),
        "missing_abc_warehouse": sum(1 for m in master if not m["ABC Indicator"]),
        "no_planning": sum(1 for m in master if m["MRP Type"] == "ND"),
        "master_without_stock": sum(1 for m in master if (m["Material"], m["Plant"]) not in stocked),
    }
    report = run_data_quality(query_service, principals["gamma.admin"])
    got = {c["id"]: c["count"] for c in report["checks"]}
    assert {k: got[k] for k in expected} == expected
    # Every planted anomaly is found.
    assert all(v > 0 for v in expected.values())
    status = {c["id"]: c["status"] for c in report["checks"]}
    assert status["no_planning"] == status["master_without_stock"] == "info"
    assert status["missing_abc_class"] == "not_mapped"  # ABC is per plant in SAP: item_warehouse.abc_class


def test_viewer_without_cost_sees_restricted_value_checks(
    query_service: QueryService, principals: dict[str, Principal]
) -> None:
    status = {c["id"]: c["status"] for c in run_data_quality(query_service, principals["gamma.viewer"])["checks"]}
    assert status["zero_valuation_price"] == status["source_value_mismatch"] == "restricted"


def test_value_maps_give_canonical_codes_and_keep_the_raw_code(
    query_service: QueryService, principals: dict[str, Principal]
) -> None:
    rows = run(
        query_service,
        principals["gamma.admin"],
        {
            "dataset": "item",
            "dimensions": ["inventory_stage", "attr_material_type"],
            "measures": [{"field": "item_id", "agg": "count", "alias": "n"}],
        },
    )
    pairs = {(r["attr_material_type"], r["inventory_stage"]) for r in rows}
    assert ("ROH", "RAW") in pairs
    assert ("ZDMO", "unmapped") in pairs
    assert {s for _, s in pairs} <= {*SAP_MATERIAL_TYPES.values(), "unmapped"}
    # Filtering works on the canonical value.
    raw = run(
        query_service,
        principals["gamma.admin"],
        {
            "dataset": "item",
            "dimensions": ["attr_material_type"],
            "filters": [{"field": "inventory_stage", "op": "eq", "value": "RAW"}],
        },
    )
    assert {r["attr_material_type"] for r in raw} == {"ROH"}


@pytest.mark.security
def test_gamma_viewer_is_scoped_to_its_plant(query_service: QueryService, principals: dict[str, Principal]) -> None:
    rows = run(
        query_service,
        principals["gamma.viewer"],
        {
            "dataset": "item_warehouse",
            "dimensions": ["warehouse_id"],
            "measures": [{"field": "item_id", "agg": "count", "alias": "n"}],
        },
    )
    assert {r["warehouse_id"] for r in rows} == {"GM01"}


# --------------------------------------------------------------------------------------- HTTP
@pytest.fixture(scope="module")
def client(data_dir: Path, tmp_path_factory: pytest.TempPathFactory) -> Iterator[TestClient]:
    settings = Settings(
        _env_file=None,  # type: ignore[call-arg]
        app_env="test",
        auth_mode="dev",
        dev_jwt_secret="gamma-secret-" + "x" * 40,
        data_dir=data_dir,
        app_db_url=app_db_url(tmp_path_factory.mktemp("gammadb") / "app.db"),
        log_level="WARNING",
    )
    with TestClient(create_app(settings)) as c:
        yield c


def login(client: TestClient, user: str) -> TestClient:
    client.cookies.clear()
    r = client.post("/api/auth/dev-login", json={"username": user}, headers=CSRF)
    assert r.status_code == 200, r.text
    return client


def test_balances_work_without_lots_abc_or_warehouse_names(client: TestClient) -> None:
    login(client, "gamma.admin")
    r = client.get("/api/inventory/balances", params={"limit": 5, "sort": "value", "dir": "desc"})
    assert r.status_code == 200, r.text
    names = [c["name"] for c in r.json()["columns"]]
    assert "stock_status" in names
    assert "stock_value" in names
    assert not {"lot_id", "item.abc_class", "warehouse.name"} & set(names)
    blocked = client.get("/api/inventory/balances", params={"status": "blocked", "limit": 50}).json()["data"]
    assert blocked
    assert {row["stock_status"] for row in blocked} == {"blocked"}
    # Alpha (with lots) still gets its full column set.
    login(client, "alpha.admin")
    names = [c["name"] for c in client.get("/api/inventory/balances", params={"limit": 1}).json()["columns"]]
    assert {"lot_id", "item.abc_class", "warehouse.name"} <= set(names)


def test_filters_list_only_what_the_source_has(client: TestClient) -> None:
    f = login(client, "gamma.admin").get("/api/inventory/filters").json()
    assert f["abc_classes"] == []
    assert "unrestricted" in f["stock_statuses"]
    assert set(f["stock_statuses"]) <= {
        "unrestricted",
        "quality_inspection",
        "in_transit",
        "restricted",
        "blocked",
        "returns",
    }
    assert {w["id"] for w in f["warehouses"]} == {"GM01", "GM02", "GM03"}
    assert "RAW" in f["stages"]


def test_capability_matrix_names_the_file_that_would_unlock_each_gap(client: TestClient) -> None:
    caps = {c["id"]: c for c in login(client, "gamma.admin").get("/api/admin/capabilities").json()["capabilities"]}
    assert caps["inventory_value"]["available"]
    assert not caps["dormant_stock"]["available"]
    assert caps["dormant_stock"]["needs_file"] == ["SAP MB51 / MSEG material movements for the last 24 months"]
    assert "MCHB" in caps["expiring_lots"]["explanation"]


def test_gamma_reports_and_schema_check(client: TestClient) -> None:
    login(client, "gamma.admin")
    report = client.get("/api/reports/value_by_warehouse").json()
    assert [c["key"] for c in report["columns"]] == ["warehouse_id", "currency", "value"]
    listing = client.get("/api/reports/stock_listing").json()
    assert "stock_status" in [c["key"] for c in listing["columns"]]
    assert client.get("/api/admin/schema").json()["verdict"] == "ok"


# --------------------------------------------------------------------------------------- customer extracts
@pytest.fixture(scope="module")
def extract_root(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """A customer_extract tenant built from twin files (synthetic values, real-extract classification)."""
    root = tmp_path_factory.mktemp("extract")
    stock, master, _ = write_twin_exports(root / "in", SAP_TWINS["demo-gamma"], as_of=AS_OF, n_materials=60)
    relative = "var/private/client-x/warehouse.duckdb"
    result = import_sap_exports(
        stock, master, tenant_id="client-x", as_of=AS_OF, target=root / relative, synthetic=False
    )
    config = sap_tenant_config(
        tenant_id="client-x",
        display_name="Client X extract",
        data_classification="customer_extract",
        source_system="Client SAP exports",
        connection_path=relative,
        plants=result.plants,
        dev_users=[{"username": "client-x.admin", "display_name": "X", "roles": ["super_admin"], "warehouses": ["*"]}],
    )
    write_tenant_yaml(root / "tenants" / "client-x.yaml", config, "test")
    return root


def extract_client(root: Path, tmp_path: Path, **overrides: Any) -> TestClient:
    settings = Settings(
        _env_file=None,  # type: ignore[call-arg]
        **{
            "app_env": "development",
            "auth_mode": "dev",
            "dev_jwt_secret": "extract-secret-" + "x" * 40,
            "data_dir": root,
            "private_tenants_dir": root / "tenants",
            "allow_customer_extracts": True,
            "app_db_url": app_db_url(tmp_path / "app.db"),
            "log_level": "WARNING",
            **overrides,
        },
    )
    return TestClient(create_app(settings))


@pytest.mark.security
def test_extract_sign_in_needs_the_flag(extract_root: Path, tmp_path: Path) -> None:
    with extract_client(extract_root, tmp_path, allow_customer_extracts=False) as c:
        users = {u["username"] for u in c.get("/api/auth/dev-users").json()["users"]}
        assert "client-x.admin" not in users
        assert "gamma.admin" in users
        r = c.post("/api/auth/dev-login", json={"username": "client-x.admin"}, headers=CSRF)
        assert r.status_code == 401


@pytest.mark.security
def test_extract_is_labelled_confidential_end_to_end(extract_root: Path, tmp_path: Path) -> None:
    with extract_client(extract_root, tmp_path) as c:
        users = {u["username"]: u for u in c.get("/api/auth/dev-users").json()["users"]}
        assert users["client-x.admin"]["data_classification"] == "customer_extract"
        assert c.post("/api/auth/dev-login", json={"username": "client-x.admin"}, headers=CSRF).status_code == 200
        meta = c.get("/api/meta/tenant").json()
        assert meta["data_classification"] == "customer_extract"
        assert meta["synthetic"] is False
        balances = c.get("/api/inventory/balances", params={"limit": 1}).json()
        assert balances["provenance"]["data_classification"] == "customer_extract"
        r = c.get("/api/reports/stock_listing/export", params={"format": "csv"})
        assert r.status_code == 200, r.text
        assert "CONFIDENTIAL-stock_listing-" in r.headers["content-disposition"]
        text = r.content.decode("utf-8")
        assert "CONFIDENTIAL - client data extract" in text
        assert "SYNTHETIC" not in text


@pytest.mark.security
def test_external_ai_stays_off_for_an_unapproved_extract(extract_root: Path, tmp_path: Path) -> None:
    with extract_client(extract_root, tmp_path, agent_provider="anthropic", anthropic_api_key="test-key") as c:
        container = c.app.state.container  # type: ignore[attr-defined]
        from lumiant.auth.principal import principal_for_dev_user

        tenant = container.tenants.get("client-x")
        provider, notice = container.agent._provider(principal_for_dev_user(tenant, tenant.dev_users[0]))
        assert provider.external is False
        assert notice is not None
        assert "not approved" in notice


@pytest.mark.security
def test_production_refuses_to_start_with_an_extract_tenant(
    extract_root: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    registry = TenantRegistry.load(Settings(_env_file=None).tenant_config_dir, extract_root / "tenants")  # type: ignore[call-arg]
    monkeypatch.setattr("lumiant.main.TenantRegistry.load", lambda *_: registry)
    settings = Settings(
        _env_file=None,  # type: ignore[call-arg]
        app_env="production",
        auth_mode="oidc",
        cookie_secure=True,
        data_dir=extract_root,
        app_db_url=app_db_url(tmp_path / "app.db"),
    )
    with pytest.raises(RuntimeError, match="cannot run in production"):
        create_app(settings)
