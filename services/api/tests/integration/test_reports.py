"""Reports and exports: authorization, metadata, injection safety, parity with the analyses."""

from __future__ import annotations

import csv
import io
from collections.abc import Iterator
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from lumiant.core.config import Settings
from lumiant.main import create_app
from lumiant.reports.service import safe_cell
from tests.helpers import raw_duckdb

pytestmark = pytest.mark.integration
CSRF = {"X-Requested-With": "lumiant"}


@pytest.fixture(scope="module")
def client(data_dir: Path, tmp_path_factory: pytest.TempPathFactory) -> Iterator[TestClient]:
    settings = Settings(
        _env_file=None,  # type: ignore[call-arg]
        app_env="test",
        auth_mode="dev",
        dev_jwt_secret="report-secret-" + "x" * 40,
        data_dir=data_dir,
        app_db_url=f"sqlite:///{(tmp_path_factory.mktemp('repdb') / 'app.db').as_posix()}",
        log_level="WARNING",
    )
    with TestClient(create_app(settings)) as c:
        yield c


def login(client: TestClient, user: str) -> TestClient:
    client.cookies.clear()
    assert client.post("/api/auth/dev-login", json={"username": user}, headers=CSRF).status_code == 200
    return client


def parse_csv(text: str) -> tuple[dict[str, str], list[list[str]]]:
    rows = list(csv.reader(io.StringIO(text.lstrip("﻿"))))
    meta = {r[0][2:]: r[1] for r in rows if r and r[0].startswith("# ")}
    body = rows[len([r for r in rows if r and r[0].startswith("# ")]) + 1 :]
    return meta, body


def test_catalog_depends_on_role(client: TestClient) -> None:
    admin = login(client, "alpha.admin").get("/api/reports").json()
    assert {r["id"] for r in admin["reports"]} >= {"value_by_warehouse", "stock_listing", "below_reorder"}
    assert admin["can_export"] is True
    analyst = login(client, "alpha.analyst").get("/api/reports").json()
    assert "value_by_warehouse" not in {r["id"] for r in analyst["reports"]}
    assert analyst["can_export"] is False


def test_preview_matches_independent_sql(client: TestClient, alpha_db: Path) -> None:
    body = login(client, "alpha.admin").get("/api/reports/value_by_warehouse").json()
    con = raw_duckdb(alpha_db)
    expected = dict(con.execute("SELECT WERKS, SUM(LABST*VERPR) FROM ZV_STOCK GROUP BY WERKS").fetchall())
    con.close()
    got = {r["warehouse_id"]: Decimal(r["value"]) for r in body["rows"]}
    assert got == {k: v for k, v in expected.items() if v is not None}
    assert body["metadata"]["synthetic"] is True
    assert any("SYNTHETIC" in n for n in body["metadata"]["notices"])


def test_csv_export_has_metadata_and_exact_values(client: TestClient) -> None:
    login(client, "alpha.admin")
    r = client.get("/api/reports/below_reorder/export", params={"format": "csv"})
    assert r.status_code == 200
    assert r.headers["content-disposition"].startswith('attachment; filename="below_reorder-')
    assert r.headers["content-type"].startswith("text/csv")
    meta, body = parse_csv(r.text)
    assert meta["Report"] == "Items below reorder level"
    assert "SYNTHETIC" in r.text
    assert meta["Access scope"] == "all warehouses in tenant"
    header, *rows = body
    assert header[:3] == ["Item", "Description", "Warehouse"]
    assert len(rows) == int(meta["Rows"].split()[0])
    on_hand = header.index("On hand")
    level = header.index("Reorder level")
    assert all(Decimal(row[on_hand]) < Decimal(row[level]) for row in rows)


def test_xlsx_export_roundtrip(client: TestClient) -> None:
    r = login(client, "alpha.admin").get("/api/reports/dormant_stock/export", params={"format": "xlsx", "days": 180})
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws is not None
    first_col = [c.value for c in ws["A"]]
    assert "Report" in first_col
    assert "days=180" in str(ws["B7"].value)
    assert "Item" in [c.value for c in ws[next(i for i, v in enumerate(first_col, 1) if v == "Item")]]


@pytest.mark.security
def test_scope_is_enforced_and_disclosed_in_exports(client: TestClient) -> None:
    r = login(client, "alpha.north_manager").get("/api/reports/stock_listing/export")
    meta, body = parse_csv(r.text)
    assert meta["Access scope"] == "3 warehouses (permission scope)"
    wh = body[0].index("Warehouse")
    assert {row[wh] for row in body[1:]} <= {"1100", "1110", "1120"}


@pytest.mark.security
def test_export_requires_permission_and_cost_columns_stay_hidden(client: TestClient) -> None:
    login(client, "alpha.analyst")
    assert client.get("/api/reports/below_reorder/export").status_code == 403
    assert client.get("/api/reports/value_by_warehouse").status_code == 403
    listing = client.get("/api/reports/stock_listing").json()
    assert not {"valuation_price", "stock_value"} & {c["key"] for c in listing["columns"]}
    assert client.get("/api/reports/nope").status_code == 404


@pytest.mark.security
def test_tenant_isolation_in_reports(client: TestClient) -> None:
    body = login(client, "beta.user").get("/api/reports/stock_listing").json()
    assert {r["warehouse_id"] for r in body["rows"]} == {"LDS1"}


@pytest.mark.security
@pytest.mark.parametrize(
    ("raw", "safe"),
    [
        ('=HYPERLINK("http://x")', '\'=HYPERLINK("http://x")'),
        ("+1+1", "'+1+1"),
        ("@SUM(A1)", "'@SUM(A1)"),
        ("-5", "-5"),
        ("-5.25", "-5.25"),
        ("plain", "plain"),
    ],
)
def test_formula_injection_is_neutralised(raw: str, safe: Any) -> None:
    assert safe_cell(raw) == safe


def test_exports_are_audited(client: TestClient, audit_store: Any) -> None:
    login(client, "alpha.admin").get("/api/reports/excess_stock/export")
    assert audit_store is not None  # the shared store is separate; the app's own store is queried below
    store = client.app.state.container.audit  # type: ignore[attr-defined]
    assert any(e.action == "export:excess_stock" and e.status == "ok" for e in store.recent("demo-alpha", limit=20))
